#!/usr/bin/env python3
"""Build the conference programme page (index.md) from Programma_synedriou.xlsx
and sync the title/authors/back-link of every abstract in abstracts/.

Usage:  python build.py [Excel file] [markdown folder]

Both arguments are optional. Without them the script uses the markdown folder
it lives in (the parent of _build) and the most recently modified
Programma_synedriou*.xlsx found in that folder or in the folders above it.

The Excel file is the single source of truth for order, times, speakers and
titles. Abstract bodies are never touched; only their front matter and the
"back to programme" link at the bottom are rewritten.

After the build, every file of the folder that differs is copied to the other folder of FOLDERS (the working copy
in the conference archive <-> the git clone of the site), so it does not matter which copy of build.py is run.
Nothing is deleted there, and a file that is newer there (edited by hand) is not overwritten: it is reported.
"""
import datetime, filecmp, json, os, re, shutil, sys
import openpyxl

# ---------------------------------------------------------------- settings
DATES = ['Πέμπτη 3 Δεκεμβρίου 2026', 'Παρασκευή 4 Δεκεμβρίου 2026', 'Σάββατο 5 Δεκεμβρίου 2026']
PAGE_TITLE = 'Πρόγραμμα συνεδρίου'
INTRO = ''   # optional sentence shown above the day buttons
ABSTRACT_URL = '{{ site.baseurl }}/abstracts/%s.html'
PROGRAMME_URL = '{{ site.baseurl }}/'          # where index.md is served

# the two copies of the site: the build writes into the one it runs from and copies the result to the other
FOLDERS = [r'D:\Users\Georgios\Desktop\Archive\Συνέδρια\2026-12-03–05 Η Ελλάδα στη ρωμαϊκή Ρωμαϊκή Οικουμένη\github',
           r'D:\Users\Georgios\Desktop\Archive\Python\GitHub\poleisinromangreece\programme']
SKIP = {'.git', '__pycache__'}                  # never copied
# where the Excel programme is looked for when no file is given (besides the folders above the site)
EXCEL_DIR = r'D:\Users\Georgios\Desktop\Archive\Συνέδρια\2026-12-03–05 Η Ελλάδα στη ρωμαϊκή Ρωμαϊκή Οικουμένη\Πρόγραμμα Συνεδρίου'

# keyword found in column D (speakers)  ->  abstract file name (without .md)
SLUGS = [
    ('Πετρόχειλος', 'Petrochilos'), ('Ευαγγελίδης', 'Evangelidis'), ('Καντηρέα', 'Kantirea'),
    ('Τσατσαρώνη', 'Tsatsaroni'), ('Αντωνιάδης', 'Antoniadis'),
    ('Παπαγιάννη', 'Papagianni'), ('Δαμάσκος', 'Damaskos-Malamidou-Papaioannou'), ('Sève', 'Seve'),
    ('Πετράκης', 'Petrakis-Tsolakis-Chaniotis'), ('Κωνσταντίνος Ζάχος', 'Zachos'), ('Ζουμπάκη', 'Zoumpaki'),
    ('Κουκοβασίλης', 'Koukovasilis'), ('Θεοδώρου', 'Theodorou'), ('Νίγδελης', 'Nigdelis'),
    ('Μητρόπουλος', 'Mitropoulos'), ('Γιαννοπούλου', 'Giannopoulou'), ('Bonanno', 'Bonanno_Aravantinos'),
    ('Πατσιαδά', 'Bairami-Patsiada'), ('Κουτρομπή', 'Koutrompi'), ('Τιβερίου', 'Tiveriou-Sverkos'),
    ('Γαλάνη', 'Galani'), ('Κουρσούμης', 'Koursoumis'), ('Goette', 'Goette'),
    ('Χριστοπούλου', 'Christopoulou'), ('Ριζάκης', 'Rizakis'), ('Τζαναβάρη', 'Tzanavari'),
    ('Di Napoli', 'Di_Napoli'), ('Καραμπίνης', 'Karambinis'), ('Κωνσταντινίδης', 'Konstantinidis'),
    ('Καζακίδη', 'Kazakidi'), ('Λεβέντη', 'Leventi'), ('Κοκκίνη', 'Kokkini-Mourthos'),
    ('Γρηγορόπουλος', 'Grigoropoulos'), ('Παπαγεωργιάδου', 'Papageorgiadou'), ('Μηντζαρίδου', 'Minartzidou'),
    ('Χρυσάνθη Τσούλη', 'Tsouli_et_al'), ('Μαρία Τσούλη', 'Tsouni'),
]
BREAKS = ('Διάλειμμα', 'Μεσημεριανό')          # rows rendered as break-row

# ---------------------------------------------------------------- formulas
class Book:
    """Tiny evaluator for the simple formulas used in the workbook."""
    REF = re.compile(r"(?:'?([^'!=+\-*/(),]+)'?!)?\$?([A-Z]{1,2})\$?(\d+)")

    def __init__(self, path):
        self.wb = openpyxl.load_workbook(path)
        self.main = self.wb.worksheets[0].title

    def value(self, sheet, coord):
        v = self.wb[sheet][coord].value
        if isinstance(v, str) and v.startswith('='):
            return self.formula(sheet, v[1:])
        # a time typed into a cell (16:30) is read as a clock time: as a fraction of the day, like the formulas
        if isinstance(v, datetime.datetime):
            v = v.time()
        if isinstance(v, datetime.time):
            return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
        if isinstance(v, datetime.timedelta):
            return v.total_seconds() / 86400
        return v

    def formula(self, sheet, f):
        if self.REF.fullmatch(f):                       # plain reference (may be text)
            m = self.REF.fullmatch(f)
            return self.value(m.group(1) or sheet, m.group(2) + m.group(3))
        f = re.sub(r'TIME\((\d+),(\d+),(\d+)\)', lambda m: str((int(m[1]) * 3600 + int(m[2]) * 60 + int(m[3])) / 86400), f)
        expr = self.REF.sub(lambda m: repr(self.value(m.group(1) or sheet, m.group(2) + m.group(3)) or 0), f)
        if not re.fullmatch(r"[\d.+\-*/() e]+", expr):
            raise ValueError(f'Unsupported formula in {sheet}: ={f}')
        return eval(expr)

def hhmm(x):
    if x is None:
        return ''
    m = round(x * 1440)
    return f'{m // 60:02d}:{m % 60:02d}'

def slug_for(speakers):
    hits = [s for k, s in SLUGS if k in speakers]
    return hits[0] if hits else None

# ---------------------------------------------------------------- parse
def parse(path):
    bk = Book(path)
    ws = bk.wb[bk.main]
    days, day, session = [], None, None
    for r in range(4, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a == 'Σύνοψη':
            break
        e = ws.cell(r, 5).value
        if isinstance(a, str) and a.startswith('='):                # day header
            day = {'label': bk.value(bk.main, f'A{r}'), 'items': []}
            days.append(day)
        elif isinstance(a, str) and a.strip():                       # session header
            day['items'].append({'type': 'session', 'title': a.strip()})
        elif isinstance(a, (int, float)):                            # talk
            sp = (ws.cell(r, 4).value or '').strip()
            day['items'].append({'type': 'talk', 'no': int(a), 'speakers': sp, 'title': (e or '').strip(),
                                 'start': hhmm(bk.value(bk.main, f'B{r}')), 'end': hhmm(bk.value(bk.main, f'C{r}')),
                                 'slug': slug_for(sp)})
        elif str(ws.cell(r, 4).value or '').strip().startswith('Πρόεδρος'):   # chair: name in column E
            day['items'].append({'type': 'chair', 'name': str(e or '').strip()})
        elif e:                                                      # break / other item
            day['items'].append({'type': 'break' if e.startswith(BREAKS) else 'event', 'title': e.strip(), 'speakers': str(ws.cell(r, 4).value or '').strip(),
                                 'start': hhmm(bk.value(bk.main, f'B{r}')), 'end': hhmm(bk.value(bk.main, f'C{r}'))})
    return days

# ---------------------------------------------------------------- render
def esc(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def short_label(label, n):
    day = DATES[n - 1].rsplit(' ', 1)[0] if n <= len(DATES) else label
    return day

def render(days):
    out = ['---', 'layout: page', f'title: {json.dumps(PAGE_TITLE, ensure_ascii=False)}', 'permalink: /', 'published: true', '---', '',
           '<!-- Generated by _build/build.py from the Excel programme. Edit the Excel and rebuild; do not edit by hand. -->', '',
           *([f'<p class="intro">{INTRO}</p>', ''] if INTRO else []),
           '<nav class="day-nav">' + ''.join(f'<a href="#day-{n}">{esc(short_label(d["label"], n))}</a>' for n, d in enumerate(days, 1)) + '</nav>', '']
    for n, d in enumerate(days, 1):
        label = d['label']
        if n <= len(DATES):
            label = f"{label.split('–')[0].strip()} – {DATES[n - 1]}"
        out += [f'<h2 id="day-{n}">{esc(label)}</h2>', '<table style="width: 100%; border-collapse: collapse">',
                '  <tr>', '    <th style="width: 10%"></th>', '    <th style="width: 10%"></th>',
                '    <th style="width: 30%"></th>', '    <th style="width: 50%"></th>', '  </tr>', '']
        items = d['items']
        for i, it in enumerate(items):
            nxt = items[i + 1] if i + 1 < len(items) else None
            if it['type'] == 'session':
                out += ['  <tr class="session-start">', f'    <td colspan="4"><b>{esc(it["title"])}</b></td>', '  </tr>', '']
            elif it['type'] == 'chair':
                out += ['  <tr class="chair-row">', f'    <td colspan="4">Πρόεδρος: {esc(it["name"])}</td>', '  </tr>', '']
            elif it['type'] == 'talk':
                cls = 'session-paper' + (' session-end' if not nxt or nxt['type'] != 'talk' else '')
                title = esc(it['title'])
                cell = f'<a href="{ABSTRACT_URL % it["slug"]}">{title}</a>' if it['slug'] else title
                out += [f'  <tr class="{cls}">', f'    <td>{it["start"]}</td>', f'    <td>{it["end"]}</td>',
                        f'    <td>{esc(it["speakers"])}</td>', f'    <td>{cell}</td>', '  </tr>', '']
            else:
                cls = 'break-row' if it['type'] == 'break' else 'independent-session'
                out += [f'  <tr class="{cls}">', f'    <td>{it["start"]}</td>', f'    <td>{it["end"]}</td>',
                        f'    <td>{esc(it.get("speakers", ""))}</td>', f'    <td>{esc(it["title"])}</td>', '  </tr>', '']
        out += ['</table>', '']
    return '\n'.join(out)

# ---------------------------------------------------------------- abstracts
BACK = re.compile(r'\n---\s*\n\[Επιστροφή στο πρόγραμμα[^\n]*\n?$')

def sync_abstract(path, talk, day_no):
    txt = open(path, encoding='utf-8').read()
    m = re.match(r'---\n(.*?)\n---\n', txt, re.S)
    fm, body = m.group(1), txt[m.end():]
    fm = re.sub(r'^title:.*$', 'title: ' + json.dumps(talk['title'], ensure_ascii=False), fm, flags=re.M)
    aff = re.search(r'^affiliation:\s*(.*)$', fm, re.M)
    aff = json.loads(aff.group(1)) if aff and aff.group(1).strip() else ''
    authors = talk['speakers']            # affiliation is shown separately by the layout
    fm = re.sub(r'^authors:.*$', 'authors: ' + json.dumps(authors, ensure_ascii=False), fm, flags=re.M)
    body = BACK.sub('', body.rstrip('\n') + '\n').rstrip('\n')
    body += f'\n\n---\n[Επιστροφή στο πρόγραμμα της {day_no}ης ημέρας]({PROGRAMME_URL}#day-{day_no})\n'
    new = f'---\n{fm}\n---\n{body}'
    if new != txt:
        open(path, 'w', encoding='utf-8', newline='\n').write(new)
        return True
    return False

def main(xlsx, folder):
    days = parse(xlsx)
    open(os.path.join(folder, 'index.md'), 'w', encoding='utf-8', newline='\n').write(render(days))
    absdir = os.path.join(folder, 'abstracts')
    used, changed, problems = set(), [], []
    for n, d in enumerate(days, 1):
        for t in d['items']:
            if t['type'] != 'talk':
                continue
            if not t['slug']:
                problems.append(f'No abstract mapping for "{t["speakers"]}" (title shown without link); add it to SLUGS.')
                continue
            if t['slug'] in used:
                problems.append(f'"{t["speakers"]}" appears more than once in the Excel.')
            used.add(t['slug'])
            p = os.path.join(absdir, t['slug'] + '.md')
            if not os.path.exists(p):
                problems.append(f'Missing file abstracts/{t["slug"]}.md for "{t["speakers"]}".')
                continue
            if sync_abstract(p, t, n):
                changed.append(t['slug'])
    orphan = sorted(f[:-3] for f in os.listdir(absdir) if f.endswith('.md') and f[:-3] not in used)
    talks = sum(1 for d in days for t in d['items'] if t['type'] == 'talk')
    print(f'Programme: {len(days)} days, {talks} talks -> index.md')
    print(f'Abstracts updated: {len(changed)}' + (f' ({", ".join(changed)})' if changed else ''))
    if orphan:
        print('Abstract files not in the programme (left untouched, not linked):', ', '.join(orphan))
    for p in problems:
        print('WARNING:', p)

def mirror(src, dst):
    """copy every file of src that differs to dst (same relative path); nothing is deleted in dst, and a file
    that is newer in dst than in src (edited there by hand) is left alone and reported"""
    if not os.path.isdir(dst):
        print(f'WARNING: {dst} not found: nothing copied there.')
        return
    copied, kept = [], []
    for root, dirs, files in os.walk(src):
        dirs[:] = [d for d in dirs if d not in SKIP]
        for f in files:
            a = os.path.join(root, f)
            rel = os.path.relpath(a, src)
            b = os.path.join(dst, rel)
            if os.path.exists(b) and filecmp.cmp(a, b, shallow=False):
                continue
            if os.path.exists(b) and os.path.getmtime(b) > os.path.getmtime(a):
                kept.append(rel)
                continue
            os.makedirs(os.path.dirname(b), exist_ok=True)
            shutil.copy2(a, b)
            copied.append(rel)
    print(f'Copied to {dst}: {len(copied)} file(s)' + (f' ({", ".join(copied)})' if copied else ''))
    for rel in kept:
        print(f'WARNING: {rel} is newer in {dst} (edited there?): not overwritten. Copy it back by hand if that version is right.')

def find_excel(start):
    d = os.path.abspath(start)
    for _ in range(4):                     # markdown, Abstracts, conference folder, ...
        hits = [os.path.join(d, f) for f in os.listdir(d)
                if f.lower().startswith('programma_synedriou') and f.lower().endswith('.xlsx') and not f.startswith('~$')]
        if hits:
            return max(hits, key=os.path.getmtime)
        d = os.path.dirname(d)
    if os.path.isdir(EXCEL_DIR):           # e.g. when run from the git clone of the site
        hits = [os.path.join(EXCEL_DIR, f) for f in os.listdir(EXCEL_DIR)
                if f.lower().startswith('programma_synedriou') and f.lower().endswith('.xlsx') and not f.startswith('~$')]
        if hits:
            return max(hits, key=os.path.getmtime)
    return None

if __name__ == '__main__':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
    here = os.path.dirname(os.path.abspath(__file__))
    folder = sys.argv[2] if len(sys.argv) > 2 else os.path.dirname(here)
    xlsx = sys.argv[1] if len(sys.argv) > 1 else find_excel(folder)
    if not xlsx or not os.path.exists(xlsx):
        sys.exit('Excel file not found. Run:  python build.py "path\\to\\Programma_synedriou_v2.xlsx"')
    print('Excel:', xlsx)
    print('Output folder:', folder)
    main(xlsx, folder)
    same = lambda p, q: os.path.normcase(os.path.abspath(p)) == os.path.normcase(os.path.abspath(q))
    if any(same(folder, f) for f in FOLDERS):        # an output folder given by hand is not mirrored
        for other in FOLDERS:
            if not same(folder, other):
                mirror(folder, other)
